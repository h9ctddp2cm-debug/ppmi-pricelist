"""Parse the PPMI Excel list into data/seed.json for the website.

Usage: python3 scripts/parse_excel.py <path-to-xlsx>
"""
import json
import re
import sys
from collections import Counter, defaultdict

import openpyxl

GREEN = "FF00B050"

SUPPLIER_ALIASES = {
    "healthy living": "Healthy Living",
    "medwell": "Medwell",
    "emp co. ltd.": "EMP Co. Ltd.",
}


def clean(v):
    if v is None:
        return ""
    s = str(v).replace("_x000D_", "").replace("\xa0", " ").strip()
    return re.sub(r"[ \t]+", " ", s)


def canon_supplier(name):
    n = clean(name)
    key = n.lower().rstrip(".").strip()
    for k, v in SUPPLIER_ALIASES.items():
        if key == k.rstrip(".").strip():
            return v
    return n


def parse_price(text):
    """Return a numeric price if the text is a single number, else None."""
    s = clean(text).replace("HK$", "").replace("$", "").replace(",", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def main(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["Sheet1"]
    title = clean(ws["A1"].value)

    team = None
    category = None
    subcategory = None
    items = []
    categories = []  # ordered (team, category) pairs
    seen_cat = set()
    subcats = defaultdict(list)  # category -> ordered subcategories
    order = 0

    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=12):
        a, b, c, d, e, f, g, h, i, j, k, l = row
        bval = clean(b.value)
        if clean(a.value):
            team = clean(a.value).replace(" Team", "")
        if not bval:
            continue
        if bval == "Items":  # header row
            continue
        fill = b.fill.fgColor.rgb if b.fill and b.fill.fill_type else None
        is_label = not clean(c.value) and not clean(d.value) and h.value is None
        if is_label and fill == GREEN:
            category = bval
            subcategory = None
            if category not in seen_cat:
                seen_cat.add(category)
                categories.append({"team": team, "name": category})
            continue
        if is_label:
            subcategory = bval
            if subcategory not in subcats[category]:
                subcats[category].append(subcategory)
            continue

        order += 1
        price_text = clean(h.value)
        if price_text.replace(".", "", 1).isdigit():
            price_text = str(int(float(price_text))) if float(price_text).is_integer() else price_text
        items.append(
            {
                "order": order,
                "team": team,
                "category": category,
                "subcategory": subcategory,
                "name": bval,
                "model": clean(c.value),
                "supplier": canon_supplier(d.value),
                "spec": clean(e.value),
                "weight": clean(f.value),
                "weight_limit": clean(g.value),
                "price_text": price_text,
                "price": parse_price(h.value),
                "sales": clean(i.value),
                "tel": clean(j.value),
                "remarks": clean(k.value),
                "url": clean(l.value),
            }
        )

    # Supplier-level default contact: most common sales / tel per supplier
    suppliers = {}
    for name in sorted({it["supplier"] for it in items}):
        rows = [it for it in items if it["supplier"] == name]
        sales = Counter(r["sales"] for r in rows if r["sales"]).most_common(1)
        tel = Counter(r["tel"] for r in rows if r["tel"]).most_common(1)
        site = Counter(
            re.match(r"https?://[^/]+", r["url"]).group(0) for r in rows if r["url"].startswith("http")
        ).most_common(1)
        email = Counter(r["remarks"] for r in rows if "@" in r["remarks"]).most_common(1)
        suppliers[name] = {
            "name": name,
            "contact_name": sales[0][0] if sales else "",
            "tel": tel[0][0] if tel else "",
            "email": email[0][0] if email else "",
            "website": site[0][0] if site else "",
            "item_count": len(rows),
        }

    out = {
        "title": title,
        "categories": categories,
        "subcategories": subcats,
        "suppliers": list(suppliers.values()),
        "items": items,
    }
    with open("data/seed.json", "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
    print(f"{len(items)} items, {len(suppliers)} suppliers, {len(categories)} categories")
    for s in out["suppliers"]:
        print(f"  {s['name']:<28} {s['item_count']:>3}  {s['contact_name']} {s['tel']}")


if __name__ == "__main__":
    main(sys.argv[1])
